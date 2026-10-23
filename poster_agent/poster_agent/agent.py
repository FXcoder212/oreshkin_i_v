"""ИИ-агент: план мероприятий -> промпты для афиш -> картинки air.fail -> афиши в рамке.

Claude читает план мероприятий и рамку, для каждого мероприятия пишет промпт
и вызывает инструмент create_poster. Инструмент генерирует картинку через
air.fail (gptimage), вставляет её в рамку и возвращает Claude превью готовой
афиши — агент проверяет результат и при необходимости перегенерирует.
"""

from __future__ import annotations

import base64
import io
import json
import re
import threading
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import anthropic
from anthropic import beta_tool
from anthropic.lib.tools import ToolError
from PIL import Image, ImageDraw

from .airfail import AirFailClient, AirFailError, closest_size
from .frame import Frame, compose

MODEL = "claude-opus-5-5"
MAX_ATTEMPTS_PER_EVENT = 3

SYSTEM_PROMPT = """\
Ты — арт-директор, который готовит афиши мероприятий. Тебе дают план мероприятий \
и фирменную рамку, в окно которой вставляется сгенерированная картинка.

Для каждого мероприятия из плана:
1. Выдели суть: название, дату/время, место, формат, аудиторию, настроение.
2. Напиши промпт для модели генерации изображений GPT Image. Промпт пиши на английском, \
подробно: сюжет и главный объект, композиция, стиль (иллюстрация/фото/3D/плакатная графика), \
палитра, свет, настроение. Согласуй палитру и стиль с рамкой, которую видишь на изображении, \
и делай афиши одной серии узнаваемыми, но разными по сюжету.
3. Учитывай, что картинка будет обрезана под окно рамки с указанным соотношением сторон: \
важное держи ближе к центру, по краям оставляй «воздух».
4. {text_rule}
5. Вызови инструмент create_poster. Посмотри на результат: если текст искажён, \
главный объект обрезан, стиль не подходит к рамке или есть артефакты — исправь промпт \
и вызови create_poster ещё раз для того же мероприятия (не более {max_attempts} попыток). \
Если всё хорошо — переходи к следующему мероприятию.

Мероприятия обрабатывай по одному, по порядку. Ничего не выдумывай сверх плана: \
если в плане нет даты или места, не пиши их на афише. \
В конце коротко перечисли по-русски, какие афиши получились и где они сохранены.\
"""

TEXT_RULE_ON = (
    "Текст на афише: попроси модель аккуратно написать на изображении название мероприятия "
    "и дату/место ровно в том виде и на том языке, как в плане — дословно, в кавычках, "
    "с указанием шрифта, размера и расположения (крупный заголовок, читаемый мелкий текст). "
    "Текста должно быть мало — только самое важное."
)
TEXT_RULE_OFF = (
    "Текст на афише не нужен: явно добавь в промпт «no text, no letters, no captions, no logos» — "
    "надписи будут наложены отдельно."
)


@dataclass
class PosterResult:
    event_id: str
    title: str
    prompt: str
    attempt: int
    poster_path: str | None
    art_path: str | None
    error: str | None = None


@dataclass
class PosterAgent:
    frame: Frame
    out_dir: Path
    image_client: AirFailClient | None  # None -> «сухой прогон»: только промпты
    with_text: bool = True
    effort: str = "medium"
    log: Callable[[str], None] = print
    on_poster: Callable[[PosterResult], None] | None = None
    stop_event: threading.Event = field(default_factory=threading.Event)
    results: list[PosterResult] = field(default_factory=list)

    def __post_init__(self):
        self.out_dir = Path(self.out_dir)
        self.out_dir.mkdir(parents=True, exist_ok=True)
        self.size = closest_size(self.frame.aspect)
        self._attempts: dict[str, int] = {}

    # ---------- инструменты для Claude ----------

    def _make_tools(self):
        agent = self

        @beta_tool
        def create_poster(event_id: str, title: str, prompt: str):
            """Сгенерировать картинку по промпту (GPT Image через air.fail), вставить её в рамку и сохранить афишу.

            Возвращает превью готовой афиши для проверки.

            Args:
                event_id: Короткий латинский идентификатор мероприятия, например "01-open-day". \
Одинаковый для всех попыток одного мероприятия.
                title: Название мероприятия как в плане.
                prompt: Полный промпт для генерации изображения (на английском).
            """
            return agent._create_poster(event_id, title, prompt)

        return [create_poster]

    def _create_poster(self, event_id: str, title: str, prompt: str):
        if self.stop_event.is_set():
            raise ToolError("Пользователь остановил работу. Больше ничего не генерируй, коротко подведи итог.")
        slug = re.sub(r"[^a-zA-Z0-9_-]+", "-", event_id).strip("-")[:60] or "event"
        attempt = self._attempts.get(slug, 0) + 1
        if attempt > MAX_ATTEMPTS_PER_EVENT:
            raise ToolError(f"Для «{title}» исчерпан лимит попыток ({MAX_ATTEMPTS_PER_EVENT}). "
                            "Оставь лучший вариант и переходи к следующему мероприятию.")
        self._attempts[slug] = attempt
        self.log(f"[{slug}] попытка {attempt}: {title}")

        stem = f"{slug}_v{attempt}"
        (self.out_dir / f"{stem}_prompt.txt").write_text(prompt, encoding="utf-8")

        if self.image_client is None:
            res = PosterResult(slug, title, prompt, attempt, None, None)
            self._record(res)
            return f"Сухой прогон: промпт сохранён в {stem}_prompt.txt, изображение не генерировалось. " \
                   "Не перегенерируй это мероприятие, переходи к следующему."

        try:
            art = self.image_client.generate(prompt, self.size)
        except AirFailError as e:
            self._record(PosterResult(slug, title, prompt, attempt, None, None, str(e)))
            raise ToolError(f"Ошибка генерации: {e}")

        art_path = self.out_dir / f"{stem}_art.png"
        poster_path = self.out_dir / f"{stem}_poster.png"
        art.save(art_path)
        poster = compose(self.frame, art)
        poster.save(poster_path)
        res = PosterResult(slug, title, prompt, attempt, str(poster_path), str(art_path))
        self._record(res)
        self.log(f"[{slug}] сохранено: {poster_path}")

        return [
            {"type": "text", "text": f"Афиша сохранена: {poster_path.name} (попытка {attempt} "
                                     f"из {MAX_ATTEMPTS_PER_EVENT}). Проверь превью."},
            _image_block(poster),
        ]

    def _record(self, res: PosterResult):
        self.results.append(res)
        (self.out_dir / "manifest.json").write_text(
            json.dumps([r.__dict__ for r in self.results], ensure_ascii=False, indent=2), encoding="utf-8")
        if self.on_poster:
            self.on_poster(res)

    # ---------- запуск ----------

    def run(self, plan: list[dict], client: anthropic.Anthropic | None = None) -> str:
        """plan — content-блоки пользовательского сообщения с планом (см. load_plan)."""
        client = client or anthropic.Anthropic()
        x, y, w, h = self.frame.box
        intro = (
            f"Рамка ниже. Окно под картинку: {w}×{h} px (соотношение {w / h:.2f}), "
            f"картинка генерируется в {self.size[0]}×{self.size[1]} и обрезается по центру под окно. "
            "Прозрачная область рамки показана шахматкой."
        )
        content = [{"type": "text", "text": intro}, _image_block(self.frame.image),
                   {"type": "text", "text": "План мероприятий:"}, *plan,
                   {"type": "text", "text": "Сделай афиши для всех мероприятий из плана."}]

        system = SYSTEM_PROMPT.format(text_rule=TEXT_RULE_ON if self.with_text else TEXT_RULE_OFF,
                                      max_attempts=MAX_ATTEMPTS_PER_EVENT)
        runner = client.beta.messages.tool_runner(
            model=MODEL,
            max_tokens=16000,
            system=system,
            thinking={"type": "adaptive"},
            output_config={"effort": self.effort},
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
            tools=self._make_tools(),
            messages=[{"role": "user", "content": content}],
            max_iterations=60,
        )

        final_text = ""
        for message in runner:
            if self.stop_event.is_set():
                self.log("Остановлено пользователем.")
                break
            if message.stop_reason == "refusal":
                raise RuntimeError("Модель отказалась выполнять запрос")
            text = "".join(b.text for b in message.content if b.type == "text").strip()
            if text:
                self.log(text)
                final_text = text
        return final_text


# ---------- загрузка плана ----------

def load_plan(path: str | Path) -> list[dict]:
    """Превращает файл с планом в content-блоки для Claude.

    Поддерживаются .txt/.md/.csv/.json/.tsv (как текст), .pdf (документом),
    .docx и .xlsx (если установлены python-docx / openpyxl), изображения.
    """
    path = Path(path)
    ext = path.suffix.lower()
    data = path.read_bytes()
    if ext == ".pdf":
        return [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                 "data": base64.b64encode(data).decode()}}]
    if ext in (".png", ".jpg", ".jpeg", ".webp"):
        return [_image_block(Image.open(io.BytesIO(data)))]
    if ext == ".docx":
        import docx  # python-docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            parts += [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        return [{"type": "text", "text": "\n".join(parts)}]
    if ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    rows.append(" | ".join("" if v is None else str(v) for v in row))
        return [{"type": "text", "text": "\n".join(rows)}]
    return [{"type": "text", "text": data.decode("utf-8", errors="replace")}]


# ---------- утилиты ----------

def _image_block(img: Image.Image, max_side: int = 1024) -> dict:
    img = img.copy()
    img.thumbnail((max_side, max_side))
    if img.mode == "RGBA":
        img = _checkerboard(img)
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=85)
    return {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                        "data": base64.b64encode(buf.getvalue()).decode()}}


def _checkerboard(img: Image.Image, cell: int = 16) -> Image.Image:
    """Подкладывает шахматку под прозрачные области, чтобы их было видно на превью."""
    bg = Image.new("RGBA", img.size, (255, 255, 255, 255))
    draw = ImageDraw.Draw(bg)
    for yy in range(0, img.height, cell):
        for xx in range((yy // cell) % 2 * cell, img.width, 2 * cell):
            draw.rectangle((xx, yy, xx + cell - 1, yy + cell - 1), fill=(204, 204, 204, 255))
    return Image.alpha_composite(bg, img.convert("RGBA"))
